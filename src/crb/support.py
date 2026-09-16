"""Téléchargement vérifié, tables de résultats et provenance des calculs."""

from __future__ import annotations

import hashlib
import json
import math
import platform
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd


def config() -> dict:
    """Lit les paramètres de l’expérience depuis la racine du dépôt."""
    return json.loads(Path("config/protocol.json").read_text())


def fingerprint(path: Path) -> str:
    """Calcule l’empreinte SHA-256 du contenu exact d’un fichier."""
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def fetch() -> None:
    """Télécharge les entrées absentes et refuse un millésime différent."""
    manifest = json.loads(Path("config/data_manifest.json").read_text())
    raw = Path("data/raw")
    raw.mkdir(parents=True, exist_ok=True)
    for name, meta in manifest.items():
        path = raw / name
        if not path.exists():
            temporary = path.with_suffix(path.suffix + ".part")
            subprocess.run(
                [
                    "curl",
                    "--fail",
                    "--location",
                    "--silent",
                    "--show-error",
                    "--retry",
                    "2",
                    "--max-time",
                    "90",
                    meta["url"],
                    "-o",
                    str(temporary),
                ],
                check=True,
            )
            temporary.replace(path)
        if fingerprint(path) != meta["sha256"]:
            raise ValueError(
                f"Le millésime de {name} diffère du manifeste. Conserver le fichier séparément et créer un nouveau manifeste avant une nouvelle étude."
            )


def check_inputs() -> None:
    """Refuse une source absente ou différente de l’empreinte publiée."""
    for name, meta in json.loads(Path("config/data_manifest.json").read_text()).items():
        path = Path("data/raw") / name
        if not path.exists() or fingerprint(path) != meta["sha256"]:
            raise ValueError(f"Source absente ou modifiée {path}. Exécuter la commande fetch.")


def save_table(frame: pd.DataFrame, name: str) -> None:
    """Écrit un tableau CSV sans index, avec douze chiffres significatifs."""
    p = Path("results/tables")
    p.mkdir(parents=True, exist_ok=True)
    frame.to_csv(p / f"{name}.csv", index=False, float_format="%.12g")


def block_indices(length: int, samples: int, horizon: int, block: int, rng) -> np.ndarray:
    """Blocs circulaires contigus, tronqués au nombre exact de périodes demandé."""
    if min(length, samples, horizon, block) < 1 or block > length:
        raise ValueError("Dimensions de rééchantillonnage invalides")
    starts = rng.integers(0, length, size=(samples, (horizon + block - 1) // block))
    return ((starts[..., None] + np.arange(block)) % length).reshape(samples, -1)[:, :horizon]


def finish_run(seconds: float) -> None:
    """Enregistre les tables et le programme exacts après réussite de l'analyse."""
    files = sorted(Path("results/tables").glob("*.csv"))
    outputs = {str(p): fingerprint(p) for p in files}
    outputs["results/workbook_data.json"] = fingerprint(Path("results/workbook_data.json"))
    sources = {str(p): fingerprint(p) for p in sorted(Path("src").rglob("*.py"))}
    Path("results/run_manifest.json").write_text(
        json.dumps(
            {
                "completed_at": datetime.now(UTC).isoformat(),
                "seconds": seconds,
                "python": sys.version,
                "platform": platform.platform(),
                "lock_sha256": fingerprint(Path("uv.lock")),
                "protocol_sha256": fingerprint(Path("config/protocol.json")),
                "data_manifest_sha256": fingerprint(Path("config/data_manifest.json")),
                "source_files": sources,
                "tables": outputs,
            },
            indent=2,
        )
        + "\n"
    )


def export_tables() -> None:
    """Entrepôt local depuis les tables effectivement calculées."""
    import duckdb

    path = Path("results")
    frames = {p.stem: pd.read_csv(p) for p in sorted((path / "tables").glob("*.csv"))}
    with duckdb.connect(str(path / "research.duckdb")) as con:
        for name, frame in frames.items():
            con.register("source_frame", frame)
            con.execute(f'CREATE OR REPLACE TABLE "{name}" AS SELECT * FROM source_frame')
            con.unregister("source_frame")


def check_results():
    """Contrôle le contrat du dernier calcul sans exiger les données brutes."""
    manifest = json.loads(Path("results/run_manifest.json").read_text())
    if fingerprint(Path("uv.lock")) != manifest["lock_sha256"]:
        raise ValueError("Le verrou des dépendances a changé")
    if fingerprint(Path("config/protocol.json")) != manifest["protocol_sha256"]:
        raise ValueError("Le protocole a changé depuis le calcul")
    if fingerprint(Path("config/data_manifest.json")) != manifest["data_manifest_sha256"]:
        raise ValueError("Le manifeste des données a changé depuis le calcul")
    current = {str(p) for p in Path("src").rglob("*.py")}
    if current != set(manifest["source_files"]):
        raise ValueError("La liste des fichiers du programme a changé")
    optional_path = Path("config/local_outputs.json")
    optional = json.loads(optional_path.read_text()) if optional_path.exists() else []
    for name, expected in {**manifest["source_files"], **manifest["tables"]}.items():
        path = Path(name)
        if not path.exists() and name in optional:
            continue
        if not path.exists() or fingerprint(path) != expected:
            raise ValueError(f"Fichier absent ou modifié depuis le calcul {name}")


def finish_publication():
    """Enregistre les empreintes des documents, figures et du classeur livrés."""
    paths = [
        Path("README.md"),
        Path("ARTICLE.md"),
        Path("docs/templates/readme.md"),
        Path("docs/templates/article.md"),
        Path("rapport/rapport.pdf"),
        Path("rapport/rapport.typ"),
        Path("results/resultats.xlsx"),
        Path("config/publication_protocol.json"),
        *sorted(Path("results/figures").glob("*")),
    ]
    Path("results/publication_manifest.json").write_text(
        json.dumps({str(p): fingerprint(p) for p in paths}, indent=2) + "\n"
    )


def verify():
    """Vérifie les empreintes et la correspondance entre Excel et les résultats."""
    check_results()
    check_publication_protocol()
    check_workbook()
    for name, expected in json.loads(Path("results/publication_manifest.json").read_text()).items():
        if fingerprint(Path(name)) != expected:
            raise ValueError(f"Publication modifiée {name}")
    source = json.loads(Path("rapport/source_manifest.json").read_text())
    if fingerprint(Path("ARTICLE.md")) != source["sha256"]:
        raise ValueError("L'article et la source du PDF ne correspondent plus")
    print("Empreintes du code, des tables publiées, des figures et de l'article vérifiées.")


def check_publication_protocol() -> None:
    """Empêche de réutiliser les conclusions relues pour un autre protocole.

    Les tables peuvent servir à une nouvelle expérience. Leur interprétation
    exige de relire les textes avant de modifier le protocole de publication.
    """
    reviewed = json.loads(Path("config/publication_protocol.json").read_text())
    current = config()
    different = sorted(k for k in reviewed.keys() | current.keys() if reviewed.get(k) != current.get(k))
    if different:
        raise ValueError(
            "Protocole différent du texte relu. Adapter les textes et graphiques avant publication. Paramètres concernés "
            + ", ".join(different)
        )


def check_workbook() -> int:
    """Compare les cellules Excel aux résultats préparés et contrôle les formules.

    La tolérance absorbe seulement l'écriture décimale JSON. Un export Excel
    ancien doit être refusé même si ses exemples arithmétiques restent exacts.
    """
    import openpyxl

    payload = json.loads(Path("results/workbook_data.json").read_text())
    data = openpyxl.load_workbook("results/resultats.xlsx", read_only=True, data_only=True)
    formulas = openpyxl.load_workbook("results/resultats.xlsx", read_only=True, data_only=False)
    count = 0

    def compare(actual, expected, location):
        nonlocal count
        equal = (
            math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-10)
            if isinstance(expected, (int, float)) and isinstance(actual, (int, float))
            else actual == expected
        )
        if not equal:
            raise ValueError(f"Classeur périmé ou modifié {location}. Régénérer l'export Excel.")
        count += 1

    try:
        for table in payload["tables"]:
            sheet = data[table["name"]]
            compare(sheet["A2"].value, payload["title"], f"{sheet.title}!A2")
            compare(sheet["A4"].value, table["subtitle"], f"{sheet.title}!A4")
            rows = [table["headers"], *table["rows"]]
            for actuals, expected in zip(
                sheet.iter_rows(min_row=7, max_row=6 + len(rows), max_col=len(table["headers"])),
                rows,
                strict=True,
            ):
                for cell, value in zip(actuals, expected, strict=True):
                    compare(cell.value, value, f"{sheet.title}!{cell.coordinate}")
        example = payload["example"]
        for cell in example["cells"]:
            address = f"B{cell['row']}"
            expected = cell.get("formula", cell.get("value"))
            compare(formulas[example["name"]][address].value, expected, f"Exemple!{address}")
        for address, expected in example["checks"].items():
            compare(data[example["name"]][address].value, expected, f"Exemple!{address}")
    finally:
        data.close()
        formulas.close()
    return count
